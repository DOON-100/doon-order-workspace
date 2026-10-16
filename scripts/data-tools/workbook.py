"""Read Excel files into private JSON or image packages without editing inputs.

Examples:
  python workbook.py inspect --input book.xlsx --output inspect.json
  python workbook.py rows --input book.xlsx --output rows.json --sheet Items --range A1:D20
  python workbook.py images --input book.xlsx --output new-image-package --sheet Items

JSON artifacts may contain private workbook content. Standard output contains
only counts, hashes, and the destination path. Existing destinations are refused.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import shutil
import sys
import tempfile
import warnings
from datetime import date, datetime, time, timedelta
from pathlib import Path

import openpyxl
from openpyxl.utils.cell import range_boundaries, get_column_letter


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def json_value(value):
    if isinstance(value, (datetime, date, time)):
        return {"type": type(value).__name__, "iso": value.isoformat()}
    if isinstance(value, timedelta):
        return {"type": "timedelta", "seconds": value.total_seconds()}
    if isinstance(value, float) and not math.isfinite(value):
        return {"type": "nonfinite-number", "value": str(value)}
    return value


def select_sheets(workbook, requested: list[str] | None):
    if not requested:
        return list(workbook.worksheets)
    if len(set(requested)) != len(requested):
        raise ValueError("Each selected worksheet must occur only once")
    if any(name not in workbook.sheetnames for name in requested):
        raise ValueError("A selected worksheet does not exist")
    return [workbook[name] for name in requested]


def bounds(sheet, selected_range: str | None):
    if not selected_range:
        result = (1, 1, max(sheet.max_column, 1), max(sheet.max_row, 1))
    else:
        try:
            result = range_boundaries(selected_range)
        except (ValueError, TypeError) as error:
            raise ValueError("Range must be a finite A1 cell range without a sheet name") from error
        if any(item is None for item in result):
            raise ValueError("Range must include finite row and column bounds")
    left, top, right, bottom = result
    if not (1 <= left <= right <= 16384 and 1 <= top <= bottom <= 1048576):
        raise ValueError("Range is outside Excel worksheet bounds")
    return result


def coordinates(box):
    left, top, right, bottom = box
    return f"{get_column_letter(left)}{top}:{get_column_letter(right)}{bottom}"


def header_values(sheet, box, header_row: int):
    if header_row < 0 or header_row > sheet.max_row:
        raise ValueError("Header row must be zero (no header) or an existing worksheet row")
    if not header_row:
        return []
    return [{"column": column, "value": json_value(sheet.cell(header_row, column).value)}
            for column in range(box[0], box[2] + 1)]


def formula_text(value):
    # Array formulas are represented by a special openpyxl object, not a string.
    if isinstance(value, str):
        return value
    if hasattr(value, "text"):
        return {"text": value.text, "ref": getattr(value, "ref", None)}
    if hasattr(value, "__iter__"):
        return dict(value)
    raise ValueError("Unsupported formula representation")


def cell_record(cell, cached):
    is_formula = cell.data_type == "f"
    return {"coordinate": cell.coordinate, "column": cell.column,
            "value": None if is_formula else json_value(cell.value),
            "formula": formula_text(cell.value) if is_formula else None,
            "cachedValue": json_value(cached.value) if is_formula else None,
            "dataType": cell.data_type, "cachedDataType": cached.data_type if is_formula else None,
            "numberFormat": cell.number_format}


def inspect_sheets(workbook, cached_workbook, selected, selected_range, header_row):
    result = []
    for sheet in selected:
        cached = cached_workbook[sheet.title]
        box = bounds(sheet, selected_range)
        formulas = errors = missing = nonempty = 0
        for row in sheet.iter_rows(min_col=box[0], min_row=box[1], max_col=box[2], max_row=box[3]):
            for cell in row:
                nonempty += cell.value is not None
                formulas += cell.data_type == "f"
                missing += cell.data_type == "f" and cached[cell.coordinate].value is None
                errors += cell.data_type == "e" or cached[cell.coordinate].data_type == "e"
        result.append({"name": sheet.title, "state": sheet.sheet_state,
                       "worksheetRows": sheet.max_row, "worksheetColumns": sheet.max_column,
                       "range": coordinates(box), "headerRow": header_row,
                       "headers": header_values(sheet, box, header_row),
                       "nonemptyCells": nonempty, "formulaCells": formulas,
                       "missingFormulaCaches": missing, "errorCells": errors,
                       "mergedRanges": [str(value) for value in sheet.merged_cells.ranges],
                       "tables": [{"name": table.name, "range": table.ref} for table in sheet.tables.values()],
                       "imageCount": len(sheet._images), "freezePanes": sheet.freeze_panes})
    return result


def extract_rows(cached_workbook, selected, selected_range, header_row, include_empty=False):
    result = []
    for sheet in selected:
        cached = cached_workbook[sheet.title]
        box = bounds(sheet, selected_range)
        headers = header_values(sheet, box, header_row)
        rows = []
        for row_index in range(max(box[1], header_row + 1), box[3] + 1):
            cells = [sheet.cell(row_index, column) for column in range(box[0], box[2] + 1)]
            if not include_empty and all(cell.value is None for cell in cells):
                continue
            rows.append({"row": row_index, "cells": [cell_record(cell, cached[cell.coordinate]) for cell in cells]})
        result.append({"name": sheet.title, "range": coordinates(box), "headerRow": header_row,
                       "headers": headers, "rows": rows})
    return result


def image_anchor(picture):
    anchor = picture.anchor
    if hasattr(anchor, "_from"):
        return {"kind": "cell", "row": anchor._from.row + 1, "column": anchor._from.col + 1,
                "rowOffsetEmu": anchor._from.rowOff, "columnOffsetEmu": anchor._from.colOff}
    if hasattr(anchor, "pos"):
        return {"kind": "absolute", "xEmu": anchor.pos.x, "yEmu": anchor.pos.y}
    raise ValueError("Unsupported image anchor")


def extract_images(selected, selected_range, directory: Path, source):
    entries = []
    images_dir = directory / "images"
    images_dir.mkdir()
    for sheet_index, sheet in enumerate(selected, start=1):
        box = bounds(sheet, selected_range)
        for picture_index, picture in enumerate(sheet._images, start=1):
            anchor = image_anchor(picture)
            if selected_range:
                if anchor["kind"] != "cell":
                    raise ValueError("A range cannot be applied to an image without a cell anchor")
                if not (box[0] <= anchor["column"] <= box[2] and box[1] <= anchor["row"] <= box[3]):
                    continue
            extension = picture.format.lower()
            if extension == "jpeg":
                extension = "jpg"
            if extension not in ("png", "jpg", "gif"):
                raise ValueError("Unsupported image format; original image bytes will not be converted")
            # openpyxl returns the embedded bytes directly for these three formats.
            data = picture._data()
            filename = f"sheet-{sheet_index:03d}-image-{picture_index:03d}.{extension}"
            path = images_dir / filename
            with path.open("xb") as stream:
                stream.write(data)
            entries.append({"sheet": sheet.title, "path": path.relative_to(directory).as_posix(),
                            "sha256": hashlib.sha256(data).hexdigest(), "bytes": len(data),
                            "width": picture.width, "height": picture.height, "anchor": anchor})
    manifest = {"schemaVersion": 1, "kind": "images", "source": source, "images": entries}
    with (directory / "manifest.json").open("x", encoding="utf-8") as stream:
        json.dump(manifest, stream, ensure_ascii=False, indent=2, allow_nan=False)
        stream.write("\n")
    return len(entries)


def run(command, input_path: Path, output_path: Path, sheets=None, selected_range=None,
        header_row=1, include_empty=False):
    if command not in ("inspect", "rows", "images"):
        raise ValueError("Unsupported workbook command")
    if output_path.is_symlink():
        raise FileExistsError("Output must not be a symbolic link")
    input_path, output_path = input_path.resolve(), output_path.resolve()
    if input_path == output_path:
        raise ValueError("Output must be different from input")
    if output_path.exists():
        raise FileExistsError("Output already exists; choose a new destination")
    if not input_path.is_file():
        raise FileNotFoundError("Input workbook is not a file")
    source_hash = sha256(input_path)
    source = {"name": input_path.name, "sha256": source_hash}
    workbook = cached = None
    # Unsupported drawing warnings must not silently turn image extraction into
    # an apparently successful incomplete package.
    with warnings.catch_warnings():
        warnings.simplefilter("error", UserWarning)
        try:
            workbook = openpyxl.load_workbook(input_path, data_only=False, keep_links=True)
            selected = select_sheets(workbook, sheets)
            if command == "images":
                output_path.parent.mkdir(parents=True, exist_ok=True)
                with tempfile.TemporaryDirectory(prefix=".workbook-images-", dir=output_path.parent) as temp:
                    staging = Path(temp)
                    count = extract_images(selected, selected_range, staging, source)
                    if sha256(input_path) != source_hash:
                        raise ValueError("Input changed while being read")
                    output_path.mkdir()  # Exclusive: never reuse an existing package.
                    try:
                        for item in staging.iterdir():
                            shutil.move(str(item), str(output_path / item.name))
                    except Exception:
                        shutil.rmtree(output_path)
                        raise
                return {"output": str(output_path), "sheetCount": len(selected), "imageCount": count,
                        "inputSha256": source_hash, "inputUnchanged": True}
            cached = openpyxl.load_workbook(input_path, data_only=True, keep_links=True)
            if command == "inspect":
                data = inspect_sheets(workbook, cached, selected, selected_range, header_row)
            else:
                data = extract_rows(cached, selected, selected_range, header_row, include_empty)
            payload = {"schemaVersion": 1, "kind": command, "source": source,
                       "formulaCachesRecalculated": False, "sheets": data}
            encoded = json.dumps(payload, ensure_ascii=False, indent=2, allow_nan=False) + "\n"
            if sha256(input_path) != source_hash:
                raise ValueError("Input changed while being read")
            output_path.parent.mkdir(parents=True, exist_ok=True)
            with output_path.open("x", encoding="utf-8") as stream:
                stream.write(encoded)
            summary = {"output": str(output_path), "sheetCount": len(data), "inputSha256": source_hash,
                       "inputUnchanged": True}
            if command == "rows":
                summary["rowCount"] = sum(len(sheet["rows"]) for sheet in data)
            return summary
        finally:
            if workbook is not None:
                workbook.close()
            if cached is not None:
                cached.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    commands = parser.add_subparsers(dest="command", required=True)
    for name in ("inspect", "rows", "images"):
        command = commands.add_parser(name)
        command.add_argument("--input", required=True, type=Path)
        command.add_argument("--output", required=True, type=Path,
                             help="New JSON file (inspect/rows) or new package directory (images)")
        command.add_argument("--sheet", action="append", help="Worksheet name, repeat to select several; default all")
        command.add_argument("--range", dest="selected_range", help="Finite A1 range; images use the anchor cell")
        if name != "images":
            command.add_argument("--header-row", type=int, default=1, help="Header row, or zero for no header")
        if name == "rows":
            command.add_argument("--include-empty", action="store_true")
    args = parser.parse_args()
    try:
        summary = run(args.command, args.input, args.output, args.sheet, args.selected_range,
                      getattr(args, "header_row", 1), getattr(args, "include_empty", False))
    except Exception as error:
        # Do not echo exception payloads: libraries can include cell text or URLs.
        print(json.dumps({"error": type(error).__name__, "message": "Workbook operation failed; verify input, selection, and unused output destination"}), file=sys.stderr)
        return 1
    print(json.dumps(summary, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
