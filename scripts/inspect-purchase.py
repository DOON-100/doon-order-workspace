import json
import sys
from pathlib import Path
from datetime import date, datetime
from collections import Counter
import openpyxl

if len(sys.argv) != 2:
    raise SystemExit('Usage: python scripts/inspect-purchase.py WORKBOOK.xlsx')
source=Path(sys.argv[1])
book=openpyxl.load_workbook(source,data_only=True,read_only=False)
def value(x):
    return x.isoformat() if isinstance(x,(date,datetime)) else x
out=[]
for ws in book.worksheets:
    rows=[[value(c.value) for c in row] for row in ws.iter_rows()]
    info={'sheet':ws.title,'rows':ws.max_row,'cols':ws.max_column,'merges':[str(m) for m in ws.merged_cells.ranges], 'firstRows':rows[:8], 'lastRows':rows[-4:]}
    out.append(info)
    if ws.max_column==14:
        details=[r for r in rows[3:] if r[0]]
        keys=Counter((r[0],r[7],r[13] or '') for r in details)
        full=Counter(tuple(r) for r in details)
        quantities={}
        for r in details: quantities[r[9]]=quantities.get(r[9],0)+(r[10] or 0)
        info['summary']={'detailRows':len(details),'purchaseOrders':len(set(r[0] for r in details)),'suppliers':len(set(r[2] for r in details)),'buyers':dict(Counter(r[3] for r in details)),'units':quantities,'duplicateKeyGroups':sum(n>1 for n in keys.values()),'duplicateKeyRows':sum(n for n in keys.values() if n>1),'identicalDuplicateRows':sum(n for n in full.values() if n>1),'blankCustomerOrderRows':sum(not r[13] for r in details),'statuses':dict(Counter(r[4] for r in details)),'closed':dict(Counter(r[5] for r in details)),'invalidRequired':[i+4 for i,r in enumerate(details) if not r[0] or not r[2] or not r[7] or r[10] is None],'duplicates':[[list(k),n] for k,n in keys.items() if n>1][:12],'blankExamples':[r for r in details if not r[13]][:4]}
Path('test-output/purchase-workbook-structure.json').write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(out,ensure_ascii=False))
